from PIL import Image,ImageOps,ImageDraw
from pathlib import Path
r=Path(r'D:\badizo-pos-main')
p=r/'outputs/01a0bd5d-f9db-73f3-b557-bdfd4dbfdc38'
im=Image.open(p/'sku-preview.png').convert('RGB');im.thumbnail((1450,1100));im.save(p/'sku-preview.jpg',quality=65)
from pypdf import PdfReader
pdf=r/'output/pdf/BADIZO_NEW_STORE_2026-09-20_TELUGU_ENGLISH.pdf'
d=PdfReader(pdf);print('PDF pages',len(d.pages))
for i,page in enumerate(d.pages):print(i+1, page.extract_text()[:70].replace('\n',' '))
import openpyxl
w=openpyxl.load_workbook(p/'NEW_SKU_UPLOAD.xlsx')
assert w.sheetnames[0]=='Upload'
s=w['Upload']
assert len(list(s.values)[0])==21
assert all(c.value is None for row in s.iter_rows(min_row=2,max_row=301) for c in row)
assert s['A2'].number_format=='@'
assert s['B2'].number_format=='@'
print('Excel: blank inputs, 21 headers, leading-zero text columns verified')
